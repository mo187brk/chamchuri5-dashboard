"""Build the data files the dashboard ships with.

Run this once (or whenever the Excel model or the load file changes):

    python scripts/build_data.py --excel Chamchuri5_Solar_Strategy_v5.xlsx \
                                 --csv hourly_clean_2019.csv

Inputs
    * hourly_clean_2019.csv  - CU-BEMS hourly load (floors x AC / light / plug).
    * the Excel strategy model - hourly PV shapes, sizing sweeps and the
      reference results used by the test-suite.

Outputs (written to ./data)
    * hourly_2019.csv.gz     - one row per hour (8,760 rows).
    * sizing_pv.csv          - PV sizing sweep snapshot (Case 4 teaser).
    * sizing_battery.csv     - battery sweep snapshot (insight card).
    * excel_reference.json   - headline results of the Excel model.

The dashboard itself never reads the Excel file: everything it needs is in ./data.
"""

from __future__ import annotations

import argparse
import json
from pathlib import Path

import numpy as np
import pandas as pd

N_HOURS = 8760
N_FLOORS = 7
SOURCE_ORIGINAL = "orig2019"


def load_hourly_csv(path: Path) -> pd.DataFrame:
    """Read the CU-BEMS file and return load by category, floor and data-quality flag."""
    raw = pd.read_csv(path, encoding="utf-8-sig", parse_dates=["datetime"])
    if len(raw) != N_HOURS:
        raise ValueError(f"Expected {N_HOURS} rows in the load file, found {len(raw)}")
    if not raw["datetime"].is_monotonic_increasing or raw["datetime"].duplicated().any():
        raise ValueError("Load timestamps must be unique and sorted")

    floors = range(1, N_FLOORS + 1)
    out = pd.DataFrame({"datetime": raw["datetime"]})
    out["load_kw"] = raw["Building_Total"]
    out["ac_kw"] = sum(raw[f"Floor{k}_AC"] for k in floors)
    out["light_kw"] = sum(raw[f"Floor{k}_Light"] for k in floors)
    out["plug_kw"] = sum(raw[f"Floor{k}_Plug"] for k in floors)
    for k in floors:
        out[f"floor{k}_kw"] = raw[f"Floor{k}_Total"]

    # Energy that was gap-filled by the dataset preparer (not an original 2019 reading).
    filled = sum(
        raw[f"Floor{k}_Total"].where(raw[f"Floor{k}_Total_src"] != SOURCE_ORIGINAL, 0.0)
        for k in floors
    )
    out["filled_kw"] = filled

    gap = (out["ac_kw"] + out["light_kw"] + out["plug_kw"] - out["load_kw"]).abs().max()
    if gap > 1e-6:
        raise ValueError(f"Categories do not add up to Building_Total (max gap {gap:.3g} kW)")
    return out


def load_pv_shapes(excel: Path, times: pd.Series) -> pd.DataFrame:
    """Hourly PV output per kWp for the existing and the new array (sheet PV_Hourly)."""
    pv = pd.read_excel(excel, sheet_name="PV_Hourly", usecols="A:E")
    pv.columns = ["datetime", "provisional", "shape_existing", "pv_existing_kw", "shape_new"]
    if len(pv) != N_HOURS or not (pv["datetime"].values == times.values).all():
        raise ValueError("PV_Hourly timestamps do not match the load file")
    if pv[["shape_existing", "shape_new"]].isna().any().any():
        raise ValueError("NaN in PV profiles (open the workbook in Excel and save it first)")
    return pv[["shape_existing", "shape_new"]].rename(
        columns={"shape_existing": "pv_shape_existing", "shape_new": "pv_shape_new"}
    )


def _find_row(ws, label: str, column: int = 1) -> int:
    """First row (1-based) whose cell in `column` equals `label`."""
    for row in range(1, ws.max_row + 1):
        if ws.cell(row=row, column=column).value == label:
            return row
    raise KeyError(f"Row labelled {label!r} not found in sheet {ws.title!r}")


def read_sizing(excel: Path) -> tuple[pd.DataFrame, pd.DataFrame]:
    """Static sizing sweeps from sheet Sizing_Sweep."""
    import openpyxl

    wb = openpyxl.load_workbook(excel, data_only=True)
    ws = wb["Sizing_Sweep"]

    head = _find_row(ws, "Added PV (kWp)")
    pv_rows = []
    for row in range(head + 1, ws.max_row + 1):
        first = ws.cell(row=row, column=1).value
        if not isinstance(first, (int, float)):
            break
        pv_rows.append([ws.cell(row=row, column=c).value for c in range(1, 13)])
    pv = pd.DataFrame(
        pv_rows,
        columns=[
            "added_kwp", "total_kwp", "saving_thb", "capex_thb", "payback_y", "npv_thb",
            "irr", "saving_zero_export_thb", "payback_zero_export_y", "npv_zero_export_thb",
            "peak_onpeak_kw", "export_kwh",
        ],
    )

    head = _find_row(ws, "Battery (kWh)")
    bat_rows = []
    for row in range(head + 1, ws.max_row + 1):
        first = ws.cell(row=row, column=1).value
        if not isinstance(first, (int, float)):
            break
        bat_rows.append([ws.cell(row=row, column=c).value for c in range(1, 9)])
    bat = pd.DataFrame(
        bat_rows,
        columns=[
            "battery_kwh", "power_kw", "best_threshold_kw", "extra_saving_thb",
            "peak_onpeak_kw", "cycles_per_year", "npv_thb", "breakeven_thb_per_kwh",
        ],
    )
    return pv, bat


def read_reference(excel: Path) -> dict:
    """Headline Excel results, used by the tests and by the methodology page."""
    import openpyxl

    wb = openpyxl.load_workbook(excel, data_only=True)
    ws = wb["Scenario_Results"]
    labels = {
        "Annual electricity bill (THB)": "bill_thb",
        "Annual saving vs S0 (THB)": "saving_thb",
        "PV production (kWh/year)": "pv_kwh",
        "PV energy exported (kWh)": "export_kwh",
        "Grid import (kWh/year)": "grid_kwh",
        "Peak on-peak demand (kW)": "peak_onpeak_kw",
        "Capex of the new equipment (THB)": "capex_thb",
        "Simple payback (years)": "payback_y",
        "NPV over horizon (THB)": "npv_thb",
        "IRR over horizon": "irr",
    }
    scenarios = ["S0", "S1", "S2", "S3"]
    ref: dict = {s: {} for s in scenarios}
    for row in range(1, ws.max_row + 1):
        key = labels.get(ws.cell(row=row, column=1).value)
        if key is None:
            continue
        for i, s in enumerate(scenarios):
            value = ws.cell(row=row, column=2 + i).value
            ref[s][key] = value if isinstance(value, (int, float)) else None

    start = wb["Start_Here"]
    ref["annual_bill_no_pv_thb"] = next(
        start.cell(row=r, column=2).value
        for r in range(1, start.max_row + 1)
        if start.cell(row=r, column=1).value == "Annual bill before any PV"
    )
    # Alternatives quoted in the Change_Log sheet (static figures, same method).
    ref["alt_south_new_array"] = {
        "pv_kwh": 131520, "saving_thb": 483005, "payback_y": 4.16, "npv_thb": 3552619,
    }
    ref["alt_new_replaces_old"] = {
        "pv_kwh": 93708, "saving_thb": 306373, "payback_y": 6.72, "npv_thb": 1460502,
    }
    return ref


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawTextHelpFormatter)
    parser.add_argument("--excel", required=True, type=Path)
    parser.add_argument("--csv", required=True, type=Path)
    parser.add_argument("--out", type=Path, default=Path(__file__).resolve().parents[1] / "data")
    args = parser.parse_args()

    args.out.mkdir(parents=True, exist_ok=True)
    hourly = load_hourly_csv(args.csv)
    shapes = load_pv_shapes(args.excel, hourly["datetime"])
    hourly = pd.concat([hourly, shapes], axis=1)
    hourly.to_csv(args.out / "hourly_2019.csv.gz", index=False, float_format="%.6f")

    pv, bat = read_sizing(args.excel)
    pv.to_csv(args.out / "sizing_pv.csv", index=False)
    bat.to_csv(args.out / "sizing_battery.csv", index=False)
    (args.out / "excel_reference.json").write_text(
        json.dumps(read_reference(args.excel), indent=2), encoding="utf-8"
    )

    energy = hourly["load_kw"].sum() / 1e3
    print(f"hourly_2019.csv.gz : {len(hourly)} rows, {energy:,.1f} MWh, peak {hourly['load_kw'].max():.1f} kW")
    print(f"sizing_pv.csv      : {len(pv)} rows | sizing_battery.csv: {len(bat)} rows")
    print(f"PV yield (kWh/kWp) : existing {np.sum(shapes['pv_shape_existing']):.1f}, new {np.sum(shapes['pv_shape_new']):.1f}")


if __name__ == "__main__":
    main()
